import csv
import hashlib
import io
import os
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile

import xlrd
from fastapi import APIRouter, Depends, File, Form, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select

from app.api.deps import DB, CurrentUser, require
from app.core.config import settings
from app.core.exceptions import BusinessError
from app.models import Attachment, Material, StockMovement
from app.schemas.domain import ImportCommit
from app.services.inventory import InventoryService

router = APIRouter(tags=["导入导出与附件"])

MATERIAL_COLUMN_ALIASES = {
    "code": ("code", "内部物料编码", "物料编码", "商品编号"),
    "name": ("name", "物料名称", "名称", "商品名称"),
    "mpn": ("mpn", "型号", "商品型号"),
    "specification": ("specification", "规格", "规格值"),
    "package": ("package", "封装", "封装规格"),
    "manufacturer": ("manufacturer", "厂家", "制造商", "品牌"),
    "supplier_part_number": ("supplier_part_number", "供应商料号", "商品编号"),
    "unit": ("unit", "单位"),
    "unit_price": ("unit_price", "单价", "商品单价(元)", "商品单价（元）"),
    "quantity": ("quantity", "库存", "初始库存", "数量", "购买数量"),
    "safety_stock": ("safety_stock", "安全库存"),
    "target_stock": ("target_stock", "目标库存"),
}
LCSC_REQUIRED_COLUMNS = {"商品编号", "名称", "商品型号", "购买数量"}
HEADER_HINTS = {
    alias.casefold().replace(" ", "")
    for aliases in MATERIAL_COLUMN_ALIASES.values()
    for alias in aliases
} | {
    value.casefold().replace(" ", "")
    for value in (
        "Comment",
        "Description",
        "Designator",
        "Footprint",
        "LibRef",
        "Quantity",
        "Reference",
        "Value",
        "Item",
        "Part Number",
        "物料描述",
        "位号",
        "用量",
        "元件名称",
    )
}


def _header(value: object) -> str:
    return str(value or "").strip()


def _has_content(row: dict) -> bool:
    return any(value is not None and str(value).strip() for value in row.values())


def _header_index(rows: list[tuple[object, ...] | list[object]]) -> int:
    """Find a real table header after optional report title/metadata rows."""

    best_index = 0
    best_score = (-1, -1)
    for index, values in enumerate(rows[:50]):
        headers = [_header(value) for value in values]
        normalized = [value.casefold().replace(" ", "") for value in headers if value]
        score = (sum(value in HEADER_HINTS for value in normalized), len(set(normalized)))
        if score > best_score:
            best_index = index
            best_score = score
    return best_index


def _rows_from_values(values: list[tuple[object, ...] | list[object]]) -> list[dict]:
    if not values:
        raise BusinessError("EMPTY_IMPORT_FILE", "文件中没有可分析的数据")
    header_index = _header_index(values)
    headers = [_header(value) for value in values[header_index]]
    if not any(headers):
        raise BusinessError("EMPTY_IMPORT_FILE", "文件中没有可用的列名")
    rows = [
        {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        for row in values[header_index + 1 :]
    ]
    return [row for row in rows if _has_content(row)]


def _first_value(row: dict, aliases: tuple[str, ...], default: object = "") -> object:
    for alias in aliases:
        value = row.get(alias)
        if value is not None and str(value).strip():
            return value
    return default


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_material_rows(rows: list[dict]) -> tuple[list[dict], str]:
    headers = {str(key).strip() for row in rows for key in row}
    is_lcsc = LCSC_REQUIRED_COLUMNS.issubset(headers)
    normalized = []
    for raw in rows:
        row = {
            field: _text(_first_value(raw, aliases, "pcs" if field == "unit" else ""))
            for field, aliases in MATERIAL_COLUMN_ALIASES.items()
        }
        row["unit"] = row["unit"] or "pcs"
        row["quantity"] = row["quantity"] or "0"
        row["unit_price"] = row["unit_price"] or "0"
        row["safety_stock"] = row["safety_stock"] or "0"
        row["target_stock"] = row["target_stock"] or "0"
        if is_lcsc:
            row["attributes"] = {
                "source": "立创商城",
                "purchase_type": _text(raw.get("购买类型")),
                "source_category": _text(raw.get("商品分类")),
                "gross_weight": _text(raw.get("单个毛重")),
                "line_amount": _text(
                    raw.get("金额(元)")
                    if raw.get("金额(元)") is not None
                    else raw.get("金额（元）")
                ),
            }
        normalized.append(row)
    return normalized, "立创商城购物车" if is_lcsc else "通用物料表"


def _row_errors(rows: list[dict]) -> list[dict]:
    errors = []
    for index, row in enumerate(rows, 2):
        if not row.get("code") or not row.get("name"):
            errors.append({"row": index, "message": "缺少物料编码或名称"})
            continue
        for field, label in (
            ("quantity", "购买数量/初始库存"),
            ("unit_price", "商品单价"),
            ("safety_stock", "安全库存"),
            ("target_stock", "目标库存"),
        ):
            try:
                value = Decimal(str(row.get(field) or 0))
            except (ValueError, ArithmeticError):
                errors.append({"row": index, "message": f"{label}不是有效数字"})
                break
            if value < 0:
                errors.append({"row": index, "message": f"{label}不能小于 0"})
                break
            if field == "quantity" and value != value.to_integral_value():
                errors.append({"row": index, "message": f"{label}必须是整数"})
                break
    return errors


def _xlsx_rows(content: bytes) -> list[dict]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return _rows_from_values(list(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()


def _xls_value(sheet, row_index: int, column_index: int, datemode: int):
    cell = sheet.cell(row_index, column_index)
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, datemode)
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat()
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    return cell.value


def _xls_rows(content: bytes) -> list[dict]:
    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        if workbook.nsheets < 1:
            raise BusinessError("EMPTY_IMPORT_FILE", "文件中没有工作表")
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows < 1:
            raise BusinessError("EMPTY_IMPORT_FILE", "文件中没有可分析的数据")
        values = [
            [_xls_value(sheet, row, column, workbook.datemode) for column in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
        return _rows_from_values(values)
    finally:
        workbook.release_resources()


def parse_upload(content: bytes, filename: str) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise BusinessError("EMPTY_IMPORT_FILE", "文件中没有可分析的数据")
            rows = [
                {str(key).strip(): value for key, value in row.items() if key and str(key).strip()}
                for row in reader
            ]
            rows = [row for row in rows if _has_content(row)]
        elif suffix == ".xlsx":
            rows = _xlsx_rows(content)
        elif suffix == ".xls":
            rows = _xls_rows(content)
        else:
            raise BusinessError("UNSUPPORTED_FILE", "仅支持 CSV、XLSX 和 XLS 文件")
    except BusinessError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        StopIteration,
        UnicodeDecodeError,
        ValueError,
        xlrd.XLRDError,
    ) as exc:
        raise BusinessError(
            "INVALID_IMPORT_FILE",
            "文件无法解析，请确认文件未损坏且扩展名与实际格式一致",
            details={"filename": filename},
        ) from exc
    if len(rows) > 5000:
        raise BusinessError("TOO_MANY_ROWS", "单次导入不能超过 5000 行")
    return rows


@router.post("/imports/materials/preview", dependencies=[Depends(require("import:manage"))])
async def import_preview(user: CurrentUser, file: UploadFile = File(...)):
    content = await file.read()
    raw_rows = parse_upload(content, file.filename or "")
    rows, detected_format = normalize_material_rows(raw_rows)
    return {
        "rows": rows,
        "total": len(rows),
        "errors": _row_errors(rows),
        "detected_format": detected_format,
    }


@router.post("/imports/materials/commit", dependencies=[Depends(require("import:manage"))])
def import_commit(p: ImportCommit, db: DB, user: CurrentUser):
    created = 0
    for row in p.rows:
        if db.scalar(select(Material.id).where(Material.code == str(row["code"]))):
            continue
        initial = Decimal(str(row.get("quantity") or 0))
        material = Material(
            code=str(row["code"]),
            name=str(row["name"]),
            mpn=str(row.get("mpn") or ""),
            specification=str(row.get("specification") or ""),
            package=str(row.get("package") or ""),
            manufacturer=str(row.get("manufacturer") or ""),
            supplier_part_number=str(row.get("supplier_part_number") or ""),
            unit=str(row.get("unit") or "pcs"),
            unit_price=Decimal(str(row.get("unit_price") or 0)),
            safety_stock=Decimal(str(row.get("safety_stock") or 0)),
            target_stock=Decimal(str(row.get("target_stock") or 0)),
            attributes=row.get("attributes") if isinstance(row.get("attributes"), dict) else {},
            created_by_id=user.id,
            updated_by_id=user.id,
        )
        db.add(material)
        db.commit()
        db.refresh(material)
        if initial > 0:
            InventoryService(db, user.id, f"import-{uuid.uuid4()}").inbound(
                material.id,
                initial,
                f"import-{uuid.uuid4()}",
                "导入初始库存",
                operation_type="initial",
            )
        created += 1
    return {"created": created, "skipped": len(p.rows) - created}


def csv_stream(rows, headers):
    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return io.BytesIO(output.getvalue().encode("utf-8"))


def _quantity_cell(value: Decimal) -> int | str:
    return int(value) if value == value.to_integral_value() else str(value.normalize())


@router.get("/exports/materials", dependencies=[Depends(require("export:view"))])
def export_materials(db: DB, user: CurrentUser):
    items = db.scalars(
        select(Material).where(Material.is_deleted.is_(False)).order_by(Material.code)
    ).all()
    body = csv_stream(
        (
            [
                x.code,
                x.name,
                x.mpn,
                x.package,
                x.unit,
                x.quantity,
                x.reserved_quantity,
                x.available_quantity,
                x.safety_stock,
                x.target_stock,
            ]
            for x in items
        ),
        [
            "code",
            "name",
            "mpn",
            "package",
            "unit",
            "quantity",
            "reserved_quantity",
            "available_quantity",
            "safety_stock",
            "target_stock",
        ],
    )
    return StreamingResponse(
        body,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=materials.csv"},
    )


@router.get("/exports/stock-movements", dependencies=[Depends(require("export:view"))])
def export_movements(db: DB, user: CurrentUser):
    items = db.scalars(select(StockMovement).order_by(StockMovement.created_at.desc())).all()
    body = csv_stream(
        (
            [
                x.movement_no,
                x.material_id,
                x.operation_type,
                _quantity_cell(x.quantity_delta),
                _quantity_cell(x.before_quantity),
                _quantity_cell(x.after_quantity),
                x.operator_id,
                x.reason,
                x.created_at,
            ]
            for x in items
        ),
        [
            "movement_no",
            "material_id",
            "operation_type",
            "quantity_delta",
            "before_quantity",
            "after_quantity",
            "operator_id",
            "reason",
            "created_at",
        ],
    )
    return StreamingResponse(
        body,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=stock_movements.csv"},
    )


@router.get("/exports/low-stock", dependencies=[Depends(require("export:view"))])
def export_low_stock(db: DB, user: CurrentUser):
    items = db.scalars(
        select(Material).where(
            Material.is_deleted.is_(False),
            Material.quantity - Material.reserved_quantity <= Material.safety_stock,
        )
    ).all()
    body = csv_stream(
        (
            [
                x.code,
                x.name,
                x.available_quantity,
                x.safety_stock,
                x.target_stock,
                max(0, x.target_stock - x.available_quantity),
            ]
            for x in items
        ),
        ["code", "name", "available", "safety_stock", "target_stock", "suggested_purchase"],
    )
    return StreamingResponse(
        body,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=low_stock.csv"},
    )


@router.post("/attachments", status_code=201, dependencies=[Depends(require("attachment:manage"))])
async def upload_attachment(
    db: DB, user: CurrentUser, file: UploadFile = File(...), material_id: int | None = Form(None)
):
    content = await file.read()
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise BusinessError("FILE_TOO_LARGE", "附件超过大小限制", 413)
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".txt", ".csv", ".xlsx"}:
        raise BusinessError("UNSAFE_FILE_TYPE", "不允许的附件类型")
    settings.attachment_dir.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}{suffix}"
    path = settings.attachment_dir / stored
    with path.open("wb") as handle:
        handle.write(content)
    item = Attachment(
        material_id=material_id,
        original_name=Path(file.filename or "file").name,
        stored_name=stored,
        mime_type=file.content_type or "application/octet-stream",
        size=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        uploaded_by_id=user.id,
    )
    db.add(item)
    db.commit()
    return {"id": item.id, "name": item.original_name, "size": item.size, "sha256": item.sha256}


@router.get("/attachments/{item_id}")
def download_attachment(item_id: int, db: DB, user: CurrentUser):
    item = db.get(Attachment, item_id)
    if not item:
        raise BusinessError("ATTACHMENT_NOT_FOUND", "附件不存在", 404)
    path = settings.attachment_dir / item.stored_name
    if not path.exists():
        raise BusinessError("ATTACHMENT_FILE_MISSING", "附件文件缺失", 404)
    return FileResponse(path, filename=item.original_name, media_type=item.mime_type)


@router.delete("/attachments/{item_id}", dependencies=[Depends(require("attachment:manage"))])
def delete_attachment(item_id: int, db: DB, user: CurrentUser):
    item = db.get(Attachment, item_id)
    if not item:
        raise BusinessError("ATTACHMENT_NOT_FOUND", "附件不存在", 404)
    path = settings.attachment_dir / item.stored_name
    if path.exists():
        os.remove(path)
    db.delete(item)
    db.commit()
    return {"message": "附件已删除"}
